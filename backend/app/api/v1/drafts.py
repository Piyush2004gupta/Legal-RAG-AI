"""
Drafts API - AI legal drafting (Petitions, Legal Notices, Bail Applications, Written Statements).
Supports editing and Word .docx export.
"""
import io
from pathlib import Path
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, FileResponse
from pydantic import BaseModel
from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor

from app.core.store import LocalDataStore, get_store
from app.core.openai_client import chat_completion

router = APIRouter()


class DraftGenerateRequest(BaseModel):
    draft_type: str  # LegalNotice | Petition | BailApplication | WrittenStatement | Reply
    title: Optional[str] = None
    court_name: Optional[str] = None
    petitioner_name: Optional[str] = None
    respondent_name: Optional[str] = None
    specific_instructions: Optional[str] = None
    prayer_details: Optional[str] = None


class DraftUpdateRequest(BaseModel):
    title: Optional[str] = None
    content: Optional[str] = None
    court_name: Optional[str] = None
    petitioner_name: Optional[str] = None
    respondent_name: Optional[str] = None
    prayer: Optional[str] = None


DRAFTING_SYSTEM_PROMPT = """You are an Indian high court advocate specializing in drafting legal pleadings, notices, and petitions.
Draft a formal, court-ready legal document adhering strictly to Indian procedural rules, conventions, and statutory language.

Structure:
1. Header & Jurisdiction (Court name / In the Matter of)
2. Memo of Parties (Petitioner / Complainant vs. Respondent / Accused)
3. Formal Title
4. Most Respectfully Showeth / Chronological Statement of Facts with exact dates
5. Grounds / Legal Provisions invoked with statutory rationale
6. Prayer / Relief sought
7. Verification clause with date and place

Do not invent facts not supported by the case context. Format with professional legal phrasing.
"""


@router.get("")
async def list_drafts(case_id: str, store: LocalDataStore = Depends(get_store)):
    drafts = store.list_drafts(case_id)
    return drafts


@router.post("/generate", status_code=201)
async def generate_draft(
    case_id: str,
    body: DraftGenerateRequest,
    store: LocalDataStore = Depends(get_store),
):
    case = store.get_case(case_id)
    if not case:
        raise HTTPException(status_code=404, detail="Case not found")

    chunks = store.get_chunks_by_case(case_id)
    snippets = []
    for c in chunks[:25]:
        doc = store.get_document(c.get("document_id"))
        doc_name = doc.get("original_filename", "Document") if doc else "Document"
        snippets.append(f"[{doc_name}, Page {c.get('page_number', 1)}]: {c.get('chunk_text', '')}")

    context = "\n\n".join(snippets)

    user_prompt = f"""
CASE TITLE: {case.get('title')}
CASE TYPE: {case.get('case_type', 'General')}
CLIENT / PETITIONER: {body.petitioner_name or case.get('client_name', 'Client')}
OPPOSITE PARTY / RESPONDENT: {body.respondent_name or case.get('opposite_party', 'Opposite Party')}
COURT / FORUM: {body.court_name or case.get('court', 'Appropriate Court of Law')}
DOCUMENT TYPE TO DRAFT: {body.draft_type}
SPECIFIC INSTRUCTIONS: {body.specific_instructions or 'Draft standard pleading based on case facts'}
PRAYER INSTRUCTIONS: {body.prayer_details or 'Standard legal relief based on grounds'}

CASE DOCUMENT EVIDENCE & FACTS:
{context}

Draft the complete legal pleading/notice now.
"""

    messages = [
        {"role": "system", "content": DRAFTING_SYSTEM_PROMPT},
        {"role": "user", "content": user_prompt},
    ]

    try:
        response = await chat_completion(
            messages=messages,
            temperature=0.2,
            max_tokens=4096,
        )
        content = response.get("content", "")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Drafting generation failed: {str(e)}")

    draft = store.create_draft({
        "case_id": case_id,
        "draft_type": body.draft_type,
        "title": body.title or f"{body.draft_type} - {case.get('title', 'Case')}",
        "content": content,
        "court_name": body.court_name or case.get("court"),
        "petitioner_name": body.petitioner_name or case.get("client_name"),
        "respondent_name": body.respondent_name or case.get("opposite_party"),
        "prayer": body.prayer_details,
    })

    return draft


@router.get("/{draft_id}")
async def get_draft(case_id: str, draft_id: str, store: LocalDataStore = Depends(get_store)):
    draft = store.get_draft(draft_id)
    if not draft or draft.get("case_id") != case_id:
        raise HTTPException(status_code=404, detail="Draft not found")
    return draft


@router.patch("/{draft_id}")
async def update_draft(
    case_id: str,
    draft_id: str,
    body: DraftUpdateRequest,
    store: LocalDataStore = Depends(get_store),
):
    updates = body.model_dump(exclude_none=True)
    draft = store.update_draft(draft_id, updates)
    if not draft or draft.get("case_id") != case_id:
        raise HTTPException(status_code=404, detail="Draft not found")
    return draft


@router.delete("/{draft_id}", status_code=204)
async def delete_draft(case_id: str, draft_id: str, store: LocalDataStore = Depends(get_store)):
    success = store.delete_draft(draft_id)
    if not success:
        raise HTTPException(status_code=404, detail="Draft not found")


@router.get("/{draft_id}/export/docx")
async def export_docx(case_id: str, draft_id: str, store: LocalDataStore = Depends(get_store)):
    draft = store.get_draft(draft_id)
    if not draft or draft.get("case_id") != case_id:
        raise HTTPException(status_code=404, detail="Draft not found")

    # Generate professional .docx
    doc = DocxDocument()

    # Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(draft.get("title", "LEGAL DOCUMENT").upper())
    title_run.bold = True
    title_run.font.size = Pt(14)
    title_run.font.name = "Times New Roman"

    # Content paragraphs
    lines = draft.get("content", "").split("\n")
    for line in lines:
        if not line.strip():
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"{draft.get('draft_type', 'LegalDraft')}_{draft_id[:8]}.docx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{draft_id}/export/pdf")
async def export_pdf(case_id: str, draft_id: str, store: LocalDataStore = Depends(get_store)):
    draft = store.get_draft(draft_id)
    if not draft or draft.get("case_id") != case_id:
        raise HTTPException(status_code=404, detail="Draft not found")

    import io
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib import colors

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'LegalTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        alignment=1, # Center
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=14
    )
    body_style = ParagraphStyle(
        'LegalBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=15,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )

    story = []
    title_text = draft.get("title", "LEGAL DRAFT").upper()
    story.append(Paragraph(title_text, title_style))
    story.append(Spacer(1, 12))

    raw_content = draft.get("content", "")
    # Clean markdown triple backticks if present
    clean_content = raw_content.replace("```plaintext", "").replace("```markdown", "").replace("```", "").strip()

    for line in clean_content.split("\n"):
        line_str = line.strip()
        if not line_str:
            story.append(Spacer(1, 6))
        else:
            # Escape HTML chars for ReportLab Paragraph
            safe_text = line_str.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_text, body_style))

    doc.build(story)
    buffer.seek(0)

    filename = f"{draft.get('draft_type', 'LegalDraft')}_{draft_id[:8]}.pdf"
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

